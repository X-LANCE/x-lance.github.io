#!/usr/bin/env python3
"""Build a structured papers spreadsheet from paper_list.txt."""

from __future__ import annotations

import argparse
import re
import time
import unicodedata
from pathlib import Path

from openpyxl import Workbook, load_workbook


SCRIPT_DIR = Path(__file__).resolve().parent
SITE_DIR = SCRIPT_DIR.parents[2]

TYPE_LABELS = {
    "journal": "Journal",
    "conference": "Conference",
    "book": "Book",
}

HEADERS = [
    "year",
    "type",
    "authors",
    "title",
    "venue_short",
    "citation",
    "files",
    "award",
    "raw",
    "links",
]

VENUE_PATTERNS = [
    ("LREC-COLING", r"\bLREC[-\s]?COLING\b"),
    ("INTERSPEECH", r"\bINTERSPEECH\b|\bInterSpeech\b|\bInterspeech\b"),
    ("ACM MM", r"\bACM\s+MM\b|\bACM\s+Multimedia\b|ACM International Conference on Multimedia"),
    ("ICASSP", r"\bICASSP\b"),
    ("NeurIPS", r"\bNeurIPS\b|Neural Information Processing Systems"),
    ("EMNLP", r"\bEMNLP\b|Empirical Methods in Natural Language Processing"),
    ("NAACL", r"\bNAACL\b|North American Chapter"),
    ("ACL", r"\bACL\b|Annual Meeting of the Association for Computational Linguistics"),
    ("COLING", r"\bCOLING\b|Computational Linguistics"),
    ("ICLR", r"\bICLR\b|International Conference on Learning Representations"),
    ("ICML", r"\bICML\b|International Conference on Machine Learning"),
    ("AAAI", r"\bAAAI\b"),
    ("COLM", r"\bCOLM\b|Conference on Language Modeling"),
    ("CVPR", r"\bCVPR\b|Computer Vision and Pattern Recognition"),
    ("ICCV", r"\bICCV\b|International Conference on Computer Vision"),
    ("ECCV", r"\bECCV\b|European Conference on Computer Vision"),
    ("ASRU", r"\bASRU\b|Automatic Speech Recognition and Understanding"),
    ("SLT", r"\bSLT\b|Spoken Language Technology"),
    ("ICME", r"\bICME\b|International Conference on Multimedia and Expo"),
    ("WSDM", r"\bWSDM\b|Web Search and Data Mining"),
    ("NLPCC", r"\bNLPCC\b|Natural Language Processing and Chinese Computing"),
    ("NCMMSC", r"\bNCMMSC\b|Man-Machine Speech Communication"),
    ("BTAS", r"\bBTAS\b|Biometrics: Theory, Applications and Systems"),
    ("TASLP", r"\bTASLP\b|Transactions on Audio, Speech, and Language Processing"),
    ("TPAMI", r"\bTPAMI\b|Pattern Analysis and Machine Intelligence"),
    ("TMM", r"\bTMM\b|Transactions on Multimedia"),
    ("JSTSP", r"\bJSTSP\b|Journal of Selected Topics in Signal Processing"),
    ("SPL", r"\bSPL\b|Signal Processing Letters"),
    ("SCIS", r"Science China Information Sciences|SCIENCE CHINA Information Sciences"),
    ("SCTS", r"Science China Technological Sciences|SCIENCE CHINA Technological Sciences"),
    ("CRPS", r"Cell Reports Physical Science"),
    ("AI Open", r"\bAI Open\b"),
]


def normalize_path(value: str | Path) -> Path:
    text = str(value)
    match = re.match(r"^([A-Za-z]):[\\/](.*)$", text)
    if match:
        drive = match.group(1).lower()
        rest = match.group(2).replace("\\", "/")
        return Path("/mnt") / drive / rest
    return Path(text).expanduser()


def clean_text(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def normalized_title(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", unicodedata.normalize("NFKC", value or "").lower())


def existing_links(path: Path) -> dict[str, str]:
    if not path.exists():
        return {}
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        headers = [clean_text(str(cell.value or "")) for cell in next(sheet.iter_rows(max_row=1))]
        title_index = headers.index("title")
        links_index = headers.index("links")
    except (KeyError, OSError, StopIteration, ValueError):
        return {}

    return {
        normalized_title(str(row[title_index] or "")): str(row[links_index] or "")
        for row in sheet.iter_rows(min_row=2, values_only=True)
        if row[title_index] and row[links_index]
    }


def restore_initials(value: str) -> str:
    return value.replace("<DOT>", ".")


def split_authors_title(citation: str) -> tuple[str, str]:
    protected = re.sub(r"\b([A-Z])\.", r"\1<DOT>", citation)
    parts = re.split(r"\.\s+", protected, maxsplit=2)
    if len(parts) < 2:
        return "", ""
    authors = restore_initials(parts[0])
    title = restore_initials(parts[1])
    return clean_text(authors), clean_text(title)


def extract_files(text: str) -> tuple[str, list[str]]:
    file_pattern = re.compile(r"(?<![A-Za-z0-9_-])([A-Za-z0-9][A-Za-z0-9._-]*?\.(?:pdf|ppt|pptx))(?![A-Za-z0-9])", re.I)
    files = [match.group(1) for match in file_pattern.finditer(text)]
    citation = file_pattern.sub("", text)
    return clean_text(citation), files


def extract_award(text: str) -> tuple[str, str]:
    match = re.search(r"\s*\(([^)]*(?:award|paper|spotlight|oral|best|outstanding)[^)]*)\)\s*$", text, re.I)
    if not match:
        return text, ""
    return clean_text(text[: match.start()]), clean_text(match.group(1))


def extract_venue_short(citation: str, fallback_type: str) -> str:
    for label, pattern in VENUE_PATTERNS:
        if re.search(pattern, citation, re.I):
            return label

    candidates = []
    for content in re.findall(r"\(([^)]{2,80})\)", citation):
        cleaned = re.sub(r"\b(?:19|20)\d{2}\b", "", content)
        tokens = re.findall(r"\b[A-Z][A-Z0-9-]{1,}\b", cleaned)
        candidates.extend(tokens)
    if candidates:
        return candidates[-1]
    return TYPE_LABELS.get(fallback_type, fallback_type.title())


def parse_paper_list(path: Path) -> list[dict[str, str]]:
    year = ""
    paper_type = ""
    papers: list[dict[str, str]] = []

    for raw_line in path.read_text(encoding="utf-8-sig").splitlines():
        line = clean_text(raw_line)
        if not line:
            continue
        if re.fullmatch(r"(?:19|20)\d{2}", line):
            year = line
            paper_type = ""
            continue
        lowered = line.lower()
        if lowered in TYPE_LABELS:
            paper_type = lowered
            continue
        if not year:
            continue

        citation_with_files, files = extract_files(line)
        citation, award = extract_award(citation_with_files)
        authors, title = split_authors_title(citation)
        venue_short = extract_venue_short(citation, paper_type or "conference")
        papers.append(
            {
                "year": year,
                "type": paper_type or "conference",
                "authors": authors,
                "title": title,
                "venue_short": venue_short,
                "citation": citation,
                "files": "; ".join(files),
                "award": award,
                "raw": line,
            }
        )
    return papers


def write_excel(rows: list[dict[str, str]], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "papers"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append([row.get(header, "") for header in HEADERS])
    widths = {
        "A": 10,
        "B": 14,
        "C": 42,
        "D": 72,
        "E": 16,
        "F": 110,
        "G": 38,
        "H": 24,
        "I": 110,
        "J": 92,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    workbook.save(path)


def build(input_path: Path, output_path: Path) -> list[dict[str, str]]:
    rows = parse_paper_list(input_path)
    link_map = existing_links(output_path)
    for row in rows:
        row["links"] = link_map.get(normalized_title(row["title"]), "")
    write_excel(rows, output_path)
    print(f"wrote {len(rows)} papers to {output_path}")
    return rows


def watch(input_path: Path, output_path: Path, interval: float) -> None:
    last_mtime = None
    while True:
        mtime = input_path.stat().st_mtime
        if mtime != last_mtime:
            build(input_path, output_path)
            last_mtime = mtime
        time.sleep(interval)


def main() -> None:
    parser = argparse.ArgumentParser(description="Parse paper_list.txt into a structured Excel file.")
    parser.add_argument("--input", default=str(SCRIPT_DIR / "paper_list.txt"), help="Path to paper_list.txt.")
    parser.add_argument("--output", default=str(SCRIPT_DIR / "papers.xlsx"), help="Output .xlsx path.")
    parser.add_argument("--watch", action="store_true", help="Rebuild whenever paper_list.txt changes.")
    parser.add_argument("--interval", type=float, default=2.0, help="Polling interval for --watch.")
    args = parser.parse_args()

    input_path = normalize_path(args.input)
    output_path = normalize_path(args.output)
    if args.watch:
        watch(input_path, output_path, args.interval)
    else:
        build(input_path, output_path)


if __name__ == "__main__":
    main()
